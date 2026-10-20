"""Tests for document processor functions."""

import io
import pytest
from pypdf import PdfWriter
from docx import Document

from utils.document_processor import (
    extract_text_from_txt,
    extract_text_from_pdf,
    extract_text_from_docx,
    sanitize_text,
    validate_text_content,
    process_uploaded_document,
    DocumentProcessingError,
    MAX_FILE_SIZE_BYTES,
)


def test_extract_text_from_txt_utf8():
    sample = "Hello, StudySloth! 🦥 This is a study note."
    raw = sample.encode("utf-8")
    assert extract_text_from_txt(raw) == sample


def test_extract_text_from_txt_latin1():
    sample = "Notes with café and résumé symbols."
    raw = sample.encode("latin-1")
    extracted = extract_text_from_txt(raw)
    assert "café" in extracted or "caf" in extracted


def test_extract_text_from_txt_empty():
    assert extract_text_from_txt(b"") == ""


def test_sanitize_text():
    raw = "Line 1\x00\n\n\n\nLine 2   \n\n\nLine 3"
    cleaned = sanitize_text(raw)
    assert "\x00" not in cleaned
    assert "Line 1" in cleaned
    assert "Line 2" in cleaned
    assert "Line 3" in cleaned
    # Ensure there are no 3 consecutive newlines
    assert "\n\n\n" not in cleaned


def test_validate_text_content():
    # Empty
    is_valid, msg = validate_text_content("")
    assert not is_valid
    assert "empty" in msg.lower()

    # Too short
    is_valid, msg = validate_text_content("Short note")
    assert not is_valid
    assert "too brief" in msg.lower()

    # Valid
    is_valid, msg = validate_text_content("This is a sufficiently long note about cellular biology and genetics.")
    assert is_valid
    assert msg == ""


def test_process_uploaded_document_txt():
    sample_text = "This is a great study note for computer science and big-o algorithms."
    filename = "test_notes.txt"
    file_bytes = sample_text.encode("utf-8")
    
    text, meta = process_uploaded_document(filename, file_bytes)
    assert text == sample_text
    assert meta["filename"] == filename
    assert meta["extension"] == "txt"
    assert meta["word_count"] == len(sample_text.split())
    assert meta["is_truncated"] is False


def test_process_uploaded_document_oversized():
    fake_bytes = b"x" * (MAX_FILE_SIZE_BYTES + 1024)
    with pytest.raises(DocumentProcessingError) as exc_info:
        process_uploaded_document("huge.txt", fake_bytes)
    assert "exceeds" in str(exc_info.value).lower()


def test_process_uploaded_document_unsupported_extension():
    with pytest.raises(DocumentProcessingError) as exc_info:
        process_uploaded_document("file.exe", b"binarycontent")
    assert "unsupported" in str(exc_info.value).lower()


def test_process_uploaded_document_docx():
    doc = Document()
    doc.add_heading("Biology 101", 0)
    doc.add_paragraph("Mitochondria is the powerhouse of the cell.")
    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    text, meta = process_uploaded_document("bio.docx", docx_bytes)
    assert "Biology 101" in text
    assert "powerhouse of the cell" in text
    assert meta["extension"] == "docx"
